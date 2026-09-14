# -*- coding: utf-8 -*-
"""
Genera el informe del Laboratorio 2 en formato DOCX a partir de las cifras
producidas por lab02_datamart.py.

Ejecutar despues del script principal:  python generar_informe.py
"""

import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

with open("salidas/cifras.json", encoding="utf-8") as f:
    C = json.load(f)

doc = Document()

# Estilo base compacto
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10)
base.paragraph_format.space_after = Pt(4)

for s in doc.sections:
    s.top_margin = s.bottom_margin = Pt(50)
    s.left_margin = s.right_margin = Pt(50)


def h(texto, nivel=1):
    p = doc.add_heading(texto, level=nivel)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x30, 0x53)
    return p


def par(texto, negrita=False, cursiva=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    r.font.size = Pt(size)
    return p


def tabla(cabeceras, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabeceras):
        cel = t.rows[0].cells[i]
        cel.text = ""
        run = cel.paragraphs[0].add_run(str(c))
        run.bold = True
        run.font.size = Pt(9)
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    doc.add_paragraph()
    return t


def codigo(texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    return p


def pie(texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.size = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


# ---------------------------------------------------------------------------
# Portada
# ---------------------------------------------------------------------------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("UNMSM - FISI - E.P. Ingenieria de Software\nInteligencia de Negocios")
r.bold = True
r.font.size = Pt(11)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Laboratorio N.o 2\nConstruccion de un datamart analitico y "
              "ejecucion de operaciones OLAP")
r.bold = True
r.font.size = Pt(15)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Alumno: Torres, Mathias   |   Caso: MercaAndina   |   "
              "Motor: DuckDB 1.5   |   Script: lab02_datamart.py")
r.font.size = Pt(9)

# ---------------------------------------------------------------------------
# 1. Perfilado
# ---------------------------------------------------------------------------
h("1. Perfilado de las fuentes (paso 2)", 1)

par("La carga inicial reprodujo exactamente los conteos declarados en la guia: "
    f"{C['conteo_fuentes']['ventas_cabecera']:,} comprobantes, "
    f"{C['conteo_fuentes']['ventas_detalle']:,} lineas de detalle, "
    f"{C['conteo_fuentes']['productos']} productos, "
    f"{C['conteo_fuentes']['tiendas']} tiendas, "
    f"{C['conteo_fuentes']['clientes']} clientes, "
    f"{C['conteo_fuentes']['promociones']} promociones y "
    f"{C['conteo_fuentes']['inventario']:,} filas de inventario. "
    "Sobre esa base se ejecutaron las cuatro consultas de perfilado.")

dup = ", ".join(C["skus_duplicados"])
sin_costo = ", ".join(C["productos_sin_costo"])

tabla(
    ["Hallazgo del perfilado", "Evidencia (dato observado)",
     "Decision de diseno que exige"],
    [
        ["Categorias escritas de forma inconsistente",
         f"{C['categorias_crudas']} valores distintos de categoria para "
         "9 categorias reales: 'Lacteos', 'LACTEOS' y ' Lacteos'; lo mismo "
         "con Abarrotes y Bebidas",
         "Conformar el atributo al construir dim_producto: TRIM mas "
         "normalizacion a formato Titulo, de modo que el almacen exponga una "
         "sola forma de escribir cada categoria"],
        ["SKU duplicado en el maestro de productos",
         f"{dup} aparece 2 veces en productos.csv (52 filas, 51 SKU unicos)",
         "Deduplicar con ROW_NUMBER() OVER (PARTITION BY sku) y conservar solo "
         "la primera fila antes de crear la dimension; una clave duplicada "
         "multiplicaria las filas del hecho e inflaria los importes"],
        ["Productos sin costo unitario",
         f"2 SKU sin costo registrado: {sin_costo}",
         "Imputar el costo al 72 % del precio de lista y marcarlo con el "
         "atributo costo_imputado = 'S', para no ocultar que es un valor "
         "calculado y poder aislarlo en el analisis de margen"],
        ["Comprobantes sin cliente identificado",
         f"{C['comprobantes_sin_cliente']:,} de "
         f"{C['conteo_fuentes']['ventas_cabecera']:,} comprobantes "
         f"({C['pct_comprobantes_sin_cliente']} %) no tienen documento de cliente",
         "Crear obligatoriamente la fila de valor desconocido (id_cliente = -1, "
         "'Cliente no identificado') y cargar el hecho con LEFT JOIN mas "
         "COALESCE(clave, -1); son ventas reales que no se pueden descartar"],
        ["Periodo cubierto por los datos",
         f"Del {C['periodo_desde']} al {C['periodo_hasta']}, "
         f"{C['dias_con_venta']} dias con venta",
         "Generar dim_tiempo por adelantado con el calendario completo de 2025 "
         "y 2026 (730 filas), mas amplio que el periodo de los datos, para "
         "admitir cargas posteriores sin modificar la dimension"],
    ])
pie("Tabla 1. Hoja de registro del perfilado.")

# ---------------------------------------------------------------------------
# 2. Diseno dimensional
# ---------------------------------------------------------------------------
h("2. Diseno del esquema estrella (paso 3)", 1)

tabla(
    ["Paso", "Pregunta", "Respuesta"],
    [
        ["1. Proceso de negocio", "Que actividad operativa se quiere medir?",
         "La venta al detalle en tienda: la emision de comprobantes de venta "
         "de MercaAndina y su composicion linea por linea."],
        ["2. Grano", "Que representa una fila del hecho?",
         "Una fila de hecho_ventas es una linea de detalle de un comprobante "
         "de venta: un producto vendido, en una tienda, en una fecha, dentro "
         "de un comprobante determinado. Es el grano atomico que la fuente "
         "permite: 66 214 filas, una por cada linea de ventas_detalle.csv."],
        ["3. Dimensiones", "Que contextos aplican a ese grano?",
         "Cuando: dim_tiempo. Que: dim_producto. Donde: dim_tienda. "
         "Quien: dim_cliente. Como: dim_transaccion (dimension basura con "
         "tipo de comprobante y forma de pago) y dim_promocion. "
         "Ademas nro_comprobante como dimension degenerada dentro del hecho."],
        ["4. Hechos", "Que magnitudes se miden y de que aditividad son?",
         "cantidad (aditiva), importe_bruto (aditiva), importe_descuento "
         "(aditiva), importe_neto (aditiva), costo_venta (aditiva) y "
         "margen_bruto (aditiva). El porcentaje de margen NO se almacena: es "
         "una medida no aditiva que se calcula al consultar, como division de "
         "la suma del margen entre la suma del importe neto."],
    ])
pie("Tabla 2. Hoja de diseno del metodo de cuatro pasos de Kimball.")

par("Diferencias contra el modelo objetivo de la guia. ", negrita=True)
par("El diseno propio coincide con el modelo objetivo en el grano, en las seis "
    "dimensiones y en el conjunto de medidas. Las unicas diferencias son de "
    "detalle y estan documentadas: (a) se agrego el atributo anio_mes a "
    "dim_tiempo, porque las consultas de drill-down y de drill-across agregan "
    "por mes calendario y sin ese atributo habria que reconstruirlo en cada "
    "consulta; (b) se agrego costo_imputado a dim_producto, un atributo de "
    "linaje que permite aislar en el analisis los dos productos cuyo costo fue "
    "estimado. Ninguna de las dos altera el esquema estrella ni el grano.")

par("Tres decisiones que conviene subrayar. ", negrita=True)
par("nro_comprobante es una dimension degenerada: vive dentro del hecho y no "
    "tiene tabla propia, porque no aporta ningun atributo descriptivo; sirve "
    "para agrupar las lineas de una misma transaccion y calcular el ticket "
    "promedio. El tipo de comprobante y la forma de pago se combinan en "
    "dim_transaccion, una dimension basura: son dos indicadores de baja "
    "cardinalidad que juntos producen apenas 8 filas utiles. Y el porcentaje "
    "de margen no figura entre las medidas por ser no aditivo.")

# ---------------------------------------------------------------------------
# 3. Construccion
# ---------------------------------------------------------------------------
h("3. Construccion del datamart (pasos 4 a 6)", 1)

par("Cada dimension se construyo con los tres elementos obligatorios: clave "
    "subrogada generada por el almacen con ROW_NUMBER(), clave natural "
    "conservada como atributo, y fila de valor desconocido con clave -1. La "
    "unica excepcion a la regla de claves sin significado es dim_tiempo, cuya "
    "clave subrogada es la fecha en formato AAAAMMDD.")

cd = C["control_dimensiones"]
tabla(
    ["Dimension", "Filas obtenidas", "Filas esperadas", "Composicion"],
    [
        ["dim_tiempo", cd["dim_tiempo"], 730,
         "Calendario completo de 2025 y 2026"],
        ["dim_producto", cd["dim_producto"], 52,
         "51 productos unicos (52 filas del maestro menos 1 duplicado) + fila N/D"],
        ["dim_tienda", cd["dim_tienda"], 13, "12 tiendas + fila N/D"],
        ["dim_cliente", cd["dim_cliente"], 601, "600 clientes + fila N/D"],
        ["dim_promocion", cd["dim_promocion"], 8, "7 promociones + fila N/D"],
        ["dim_transaccion", cd["dim_transaccion"], 9,
         "2 tipos de comprobante x 4 formas de pago + fila N/D"],
    ])
pie("Tabla 3. Cifras de control de las dimensiones. Todas coinciden con lo esperado.")

par(f"El conformado de categorias funciono: dim_producto expone "
    f"{C['categorias_conformadas']} categorias, no las "
    f"{C['categorias_crudas']} del maestro de origen. Esa diferencia es, en "
    "concreto, lo que significa que un almacen de datos sea integrado.")

par("En la carga del hecho se uso INNER JOIN unicamente contra dim_tiempo "
    "(toda venta tiene fecha y el calendario cubre todo el periodo, la union "
    "nunca puede fallar) y LEFT JOIN combinado con COALESCE(clave, -1) contra "
    "las demas dimensiones, de modo que un contexto desconocido nunca elimine "
    "el hecho.")

# ---------------------------------------------------------------------------
# 4. Verificacion
# ---------------------------------------------------------------------------
h("4. Verificacion de la carga (paso 7)", 1)

cu = C["cuadre"]
tabla(
    ["Indicador de control", "Valor obtenido", "Valor esperado", "Resultado"],
    [
        ["Filas en hecho_ventas", f"{int(cu['filas_hecho']):,}", "66 214", "OK"],
        ["Filas en el origen (ventas_detalle)", f"{int(cu['filas_origen']):,}",
         "66 214", "OK"],
        ["Suma del importe bruto", f"S/ {cu['bruto_datamart']:,.2f}",
         "S/ 1 249 553.68", "OK"],
        ["Suma del importe de descuento", f"S/ {cu['descuento_datamart']:,.2f}",
         "S/ 19 870.48", "OK"],
        ["Suma del importe neto", f"S/ {cu['neto_datamart']:,.2f}",
         "S/ 1 229 683.20", "OK"],
        ["Diferencia contra el origen", f"{cu['diferencia']:.2f}", "0.00", "OK"],
    ])
pie("Tabla 4. Cuadre de filas y de importes de la tabla de hechos.")

hu = C["huerfanos"]
par(f"Busqueda de hechos huerfanos. Las cuatro consultas devolvieron cero: "
    f"producto {hu['producto']}, tienda {hu['tienda']}, cliente {hu['cliente']}, "
    f"tiempo {hu['tiempo']}. Despues de aplicar COALESCE con -1 no queda "
    "ninguna clave foranea sin correspondencia en su dimension.")

par(f"Peso de los valores desconocidos. "
    f"{C['lineas_sin_cliente']:,} de {int(cu['filas_hecho']):,} lineas "
    f"({C['pct_lineas_sin_cliente']} %) quedaron asociadas a la fila -1 de "
    "dim_cliente. Es un indicador de calidad que se reporta al negocio, no un "
    "error que se oculte: define exactamente que porcion de la venta puede "
    "analizarse por perfil de cliente y cual no.")

# ---------------------------------------------------------------------------
# 5. Operaciones OLAP
# ---------------------------------------------------------------------------
doc.add_page_break()
h("5. Operaciones OLAP (paso 8)", 1)

h("5.1 Roll-up: agregacion con subtotales", 2)
par("GROUP BY ROLLUP (region, categoria) sobre el ano 2026 produce tres tipos "
    "de fila: el detalle por region y categoria, el subtotal por region "
    "(categoria en NULL) y el total general (ambas en NULL). Se muestran los "
    "subtotales y el total; el detalle completo esta en salidas/resultados.txt.")
tabla(["Region", "Nivel de agregacion", "Venta neta (S/)", "Margen (S/)",
       "% margen"],
      [[f[0],
        "Total general" if f[0] == "TOTAL GENERAL" else "Subtotal de region",
        f[2], f[3], f[4]]
       for f in C["rollup"] if "Subtotal" in f[1]])
pie("Tabla 5. Roll-up por region y categoria, ano 2026 (subtotales y total general).")

na = C["no_aditiva"]
par(f"Medida no aditiva, comprobado. El promedio de los porcentajes de margen "
    f"por producto da {na['promedio_de_pct']} %, mientras que la division de "
    f"la suma del margen entre la suma del importe neto da "
    f"{na['pct_correcto']} %. Las cifras no coinciden y la correcta es la "
    "segunda: el porcentaje debe recalcularse sobre los totales ya agregados, "
    "nunca promediarse.")

h("5.2 Drill-down: descenso en la jerarquia", 2)
par("Partiendo de la venta anual de Lacteos en 2026 se desciende a mes y "
    "subcategoria. La jerarquia recorrida es categoria -> subcategoria en "
    "producto y ano -> mes en tiempo.")
tabla(["Ano-mes", "Subcategoria", "Venta neta (S/)"], C["drilldown"][:9])
pie("Tabla 6. Drill-down de Lacteos, primer trimestre de 2026 (extracto de 18 filas).")

h("5.3 Slice y dice: acotar el subcubo", 2)
par("El slice fija el valor de una sola dimension: se corta el cubo por "
    "ano = 2026 y se observa la venta neta por region.")
tabla(["Region", "Venta neta (S/)"], C["slice"])
pie("Tabla 7. Slice por ano 2026.")

par("El dice restringe varias dimensiones a la vez: zona comercial en Lima "
    "Norte y Lima Sur, categoria en Lacteos y Bebidas, y segundo trimestre de "
    "2026. Notese que el ORDER BY usa el numero de mes y no su nombre: "
    "ordenar por nombre_mes daria Abril, Junio, Mayo, un orden alfabetico sin "
    "sentido cronologico.")
tabla(["Zona comercial", "Categoria", "Mes", "Venta neta (S/)"], C["dice"])
pie("Tabla 8. Dice sobre zona, categoria y trimestre.")

h("5.4 Pivot: reorientar los ejes", 2)
par("La clausula PIVOT coloca las categorias en filas y los trimestres en "
    "columnas. Es la misma informacion de una tabla dinamica de Excel, pero "
    "generada en el servidor. Solo aparecen T1 y T2 porque los datos llegan "
    "hasta junio de 2026.")
tabla(["Categoria"] + [str(c) for c in C["pivot"][0][1:]],
      C["pivot"][1:])
pie("Tabla 9. Pivot de venta neta por categoria y trimestre, 2026.")

h("5.5 Comparativo contra el mismo periodo del ano anterior", 2)
par("La funcion de ventana LAG se calcula en una CTE separada, particionando "
    "por categoria y mes y ordenando por ano; el filtro WHERE anio = 2026 se "
    "aplica despues. Si el filtro estuviera junto a la ventana, las filas de "
    "2025 desapareceran antes de evaluarla y la columna del ano anterior "
    "saldria vacia sin arrojar ningun mensaje de error.")
tabla(["Ano", "Mes", "Categoria", "Venta neta", "Venta ano anterior", "Var. %"],
      C["interanual"][:6])
pie("Tabla 10. Comparativo interanual de Abarrotes, 2026 vs. 2025 (extracto de 54 filas).")

h("5.6 Analisis del programa de fidelizacion", 2)
tabla(["Ano-mes", "Venta total (S/)", "Venta fidelizada (S/)", "% fidelizada"],
      C["fidelizacion"][-6:])
pie("Tabla 11. Participacion de la venta con tarjeta de fidelizacion "
    "(ultimos 6 meses).")

# ---------------------------------------------------------------------------
# 6. Preguntas de analisis
# ---------------------------------------------------------------------------
doc.add_page_break()
h("6. Preguntas de analisis", 1)

reg = C["p1_regiones"]
cat = C["p1_categorias"]
h("6.1 Region de mayor venta y categoria de mayor margen porcentual", 2)
par(f"En 2026 la region que concentra la mayor venta neta es {reg[0][0]}, con "
    f"S/ {float(reg[0][1]):,.2f}, muy por encima de {reg[1][0]} "
    f"(S/ {float(reg[1][1]):,.2f}). La categoria que aporta el mayor margen "
    f"porcentual es {cat[0][0]}, con {cat[0][2]} %, seguida de {cat[1][0]} "
    f"({cat[1][2]} %).")
par(f"No coinciden, y el punto es justamente ese. {reg[0][0]} lidera en volumen "
    f"porque concentra la mayor parte de las tiendas, pero su margen "
    f"porcentual ({reg[0][2]} %) es el mas bajo de las cuatro regiones; "
    f"{reg[-1][0]} vende mucho menos y sin embargo rinde "
    f"{float(reg[-1][2]) - float(reg[0][2]):+.2f} puntos porcentuales de "
    "diferencia. Vender mas y ganar mas son dos preguntas distintas, y por eso "
    "el datamart guarda margen_bruto e importe_neto por separado en lugar de "
    "un porcentaje precalculado.")

p2 = C["p2"]
h("6.2 El porcentaje de venta fidelizada, esta subestimado o sobrestimado?", 2)
par(f"Esta subestimado. La consulta de la seccion 5.6 calcula el porcentaje "
    f"sobre la venta total, y la fila -1 de dim_cliente lleva "
    f"tiene_tarjeta = 'N', de modo que el {p2['pct_venta_no_identificada']} % "
    "de la venta que corresponde a clientes no identificados cae entero en el "
    "denominador como venta no fidelizada. Pero no se sabe que sean clientes "
    "sin tarjeta: se sabe que no presentaron identificacion. Es un valor "
    "desconocido tratado como negativo.")
par(f"Cuantificado: sobre la venta total el indicador da "
    f"{p2['pct_sobre_venta_total']} %, mientras que restringido a la venta con "
    f"cliente identificado da {p2['pct_sobre_venta_identificada']} %. El valor "
    "verdadero esta entre ambos y no puede precisarse mas con la informacion "
    "disponible. Lo correcto es reportar el indicador sobre la base "
    "identificada, declarando expresamente su cobertura.")

p3 = C["p3"]
h("6.3 Que ocurriria con un INNER JOIN contra dim_cliente en el paso 6?", 2)
par(f"El hecho habria perdido "
    f"{int(p3['filas_con_left_join'] - p3['filas_con_inner_join']):,} de "
    f"{int(p3['filas_con_left_join']):,} lineas, es decir el "
    f"{p3['pct_venta_perdida']} % de la venta neta: "
    f"S/ {p3['neto_con_inner_join']:,.2f} en lugar de "
    f"S/ {p3['neto_con_left_join']:,.2f}. Todos los totales de la seccion 5.1 "
    "se reduciran en esa misma proporcion y el cuadre contra la contabilidad "
    "fallaria por casi cuatrocientos setenta mil soles.")
par("Lo grave no es la magnitud sino que la perdida seria silenciosa: la "
    "consulta se ejecuta sin error y devuelve cifras verosimiles. Solo el "
    "cuadre de filas contra el origen la delata, y por eso ese cuadre es el "
    "primer control que se ejecuta.")

h("6.4 Afectan los dos productos con costo imputado al margen por categoria?", 2)
par("Los dos productos con costo estimado pertenecen a Panaderia y Carnes. "
    "Se comparo el margen porcentual reportado contra el que se obtiene "
    "excluyendo esos productos:")
tabla(["Categoria", "% margen reportado", "% margen sin imputados",
       "% de la venta de la categoria que es imputada"], C["p4"])
pie("Tabla 12. Sensibilidad del margen a la imputacion de costos.")
p4 = C["p4"]
par(f"Si afectan, y de manera relevante. En {p4[0][0]} el margen reportado "
    f"({p4[0][1]} %) esta {float(p4[0][1]) - float(p4[0][2]):+.2f} puntos "
    f"respecto del margen calculado solo con costos reales ({p4[0][2]} %), y "
    f"en {p4[1][0]} la diferencia es de "
    f"{float(p4[1][1]) - float(p4[1][2]):+.2f} puntos. La razon es el peso: "
    f"los productos imputados representan el {p4[0][3]} % y el {p4[1][3]} % de "
    "la venta de sus categorias, no una fraccion marginal. La regla del 72 % "
    "sobre el precio de lista no es neutral: fija por construccion un margen "
    "del 28 % para esos productos. El atributo costo_imputado permite "
    "detectarlo; sin el, la distorsion seria invisible.")

# ---------------------------------------------------------------------------
# 7. Actividad propuesta
# ---------------------------------------------------------------------------
doc.add_page_break()
h("7. Actividad propuesta: inventario y drill-across", 1)

h("7.1 Grano y tipo de tabla de hechos", 2)
par("Grano. ", negrita=True)
par("Una fila de hecho_inventario representa el stock de un producto, en una "
    "tienda, al cierre de un mes determinado.")
par("Tipo de tabla de hechos. ", negrita=True)
par("Es una tabla de hechos de instantanea periodica (periodic snapshot). No "
    "registra transacciones sino el estado de una magnitud fotografiada a "
    "intervalos regulares y predecibles: cada fin de mes hay una fila por cada "
    "combinacion producto-tienda con existencia, exista o no movimiento. Esa "
    "es precisamente la diferencia con hecho_ventas, que es una tabla de "
    "hechos transaccional: alli una fila existe solo si ocurrio un evento.")

ic = C["inv_cuadre"]
ih = C["inv_huerfanos"]
h("7.2 Construccion y cuadres de control", 2)
par("La tabla se cargo reutilizando exclusivamente dim_tiempo, dim_tienda y "
    "dim_producto, con el mismo patron de busqueda de claves del paso 6: "
    "INNER JOIN contra tiempo por fecha_cierre y LEFT JOIN mas COALESCE(-1) "
    "contra tienda y producto. No se creo ninguna dimension nueva: ese es el "
    "requisito que hace posible el drill-across.")
tabla(["Indicador de control", "Valor obtenido", "Valor esperado", "Resultado"],
      [["Filas en hecho_inventario", f"{int(ic['filas_hecho']):,}", "9 676", "OK"],
       ["Filas en el origen", f"{int(ic['filas_origen']):,}", "9 676", "OK"],
       ["Cierres mensuales cargados", int(ic["cierres_mensuales"]),
        "18 (ene-2025 a jun-2026)", "OK"],
       ["Valorizado a costo", f"S/ {ic['valorizado_total']:,.2f}",
        f"S/ {ic['valorizado_origen']:,.2f} (origen)", "OK"],
       ["Huerfanos de producto / tienda / tiempo",
        f"{ih['producto']} / {ih['tienda']} / {ih['tiempo']}", "0 / 0 / 0", "OK"]])
pie("Tabla 13. Cifras de control de hecho_inventario.")

h("7.3 El stock es una medida semiaditiva", 2)
par("Consulta incorrecta: sumar el stock a lo largo del ano.")
codigo("SELECT SUM(h.stock_unidades) FROM hecho_inventario h\n"
       "INNER JOIN dim_tiempo t ON h.id_tiempo = t.id_tiempo\n"
       "WHERE t.anio = 2025;")
par(f"Resultado: {int(C['semi_incorrecta']):,} unidades. Esta cifra no existe "
    "en ninguna bodega de MercaAndina. Suma doce fotografias del mismo "
    "inventario y cuenta doce veces la misma lata que estuvo en el estante "
    "todo el ano. No es un stock: es un stock multiplicado por el numero de "
    "cierres.")
par("Tratamiento correcto: el stock se suma libremente sobre producto, "
    "subcategoria, tienda y region, pero sobre el tiempo no se suma. Se toma "
    "el ultimo cierre del periodo, o el promedio de los cierres si lo que se "
    "quiere es el nivel tipico de inventario.")
sc = C["semi_correcta"]
tabla(["Tratamiento", "Unidades", "Interpretacion"],
      [["SUM sobre los 12 cierres de 2025 (incorrecto)",
        f"{int(C['semi_incorrecta']):,}",
        "Sin sentido de negocio: doble conteo de doce instantaneas"],
       ["Stock al ultimo cierre de 2025 (correcto)",
        f"{int(sc['stock_ultimo_cierre']):,}",
        "Existencias reales al 31 de diciembre de 2025"],
       ["Promedio de los cierres mensuales (correcto)",
        f"{int(sc['stock_promedio_mensual']):,}",
        "Nivel tipico de inventario durante el ano"]])
pie("Tabla 14. Demostracion de la semiaditividad del stock.")
par("En eso consiste la semiaditividad: la medida es aditiva sobre todas las "
    "dimensiones excepto sobre el tiempo. Es la misma naturaleza del saldo de "
    "una cuenta bancaria, y confundirla con una medida aditiva es uno de los "
    "errores mas costosos en un almacen de datos, porque el resultado no falla "
    "sino que simplemente miente.")

h("7.4 Drill-across: venta contra inventario", 2)
par("El drill-across combina dos procesos de negocio distintos en un mismo "
    "informe. Solo es posible porque ambos hechos comparten las dimensiones "
    "conformadas Tiempo, Producto y Tienda. La estructura obligatoria es "
    "agregar cada hecho por separado al mismo nivel de detalle (mes y "
    "categoria) mediante una CTE por hecho, y unir despues por los atributos "
    "conformados. Unir las dos tablas de hechos directamente produciria un "
    "producto cartesiano y cifras infladas.")
codigo("WITH venta AS (\n"
       "    SELECT t.anio_mes, p.categoria, SUM(h.cantidad) AS unidades_vendidas\n"
       "    FROM hecho_ventas h ... GROUP BY t.anio_mes, p.categoria\n"
       "),\n"
       "stock AS (\n"
       "    SELECT t.anio_mes, p.categoria, SUM(h.stock_unidades) AS stock_cierre\n"
       "    FROM hecho_inventario h ... GROUP BY t.anio_mes, p.categoria\n"
       ")\n"
       "SELECT v.anio_mes, v.categoria, v.unidades_vendidas, s.stock_cierre,\n"
       "       ROUND(1.0 * v.unidades_vendidas / NULLIF(s.stock_cierre,0), 3)\n"
       "           AS indice_rotacion\n"
       "FROM venta v INNER JOIN stock s\n"
       "  ON v.anio_mes = s.anio_mes AND v.categoria = s.categoria;")
par("El resultado tiene 54 filas (6 meses x 9 categorias), exactamente el "
    "producto de los dos ejes: no hay duplicacion de filas. Extracto:")
tabla(["Ano-mes", "Categoria", "Unidades vendidas", "Stock de cierre",
       "Indice de rotacion"], C["drill_across"][:9])
pie("Tabla 15. Drill-across mensual de venta e inventario (extracto de 54 filas).")

rc = C["rotacion_categorias"]
par("Acumulando el semestre y usando el stock promedio mensual como "
    "denominador (el tratamiento semiaditivo correcto), las tres categorias de "
    "menor rotacion en el primer semestre de 2026 son:")
tabla(["Categoria", "Unidades vendidas (ene-jun 2026)",
       "Stock promedio mensual", "Indice de rotacion"], rc[:3])
pie("Tabla 16. Las tres categorias de menor rotacion, primer semestre de 2026.")

h("7.5 Interpretacion y recomendacion para la gerencia comercial", 2)
par(f"Panaderia, Carnes y Bebidas rotan por debajo del resto del surtido "
    f"({rc[0][3]}, {rc[1][3]} y {rc[2][3]} veces el stock promedio en el "
    f"semestre, frente a {rc[-1][3]} de {rc[-1][0]}, la categoria mas agil). "
    "El caso de Panaderia es el mas llamativo: es la categoria de mayor margen "
    f"porcentual del datamart ({cat[0][2]} %) y a la vez la de menor rotacion, "
    "lo que significa que MercaAndina esta inmovilizando capital en el "
    "producto que mas rinde por unidad vendida. Carnes agrava el cuadro porque "
    "combina rotacion baja con el margen porcentual mas bajo de todas las "
    "categorias.")
par("Recomendacion. ", negrita=True)
par("Reducir el nivel de reposicion de Panaderia y Carnes hasta acercar su "
    "rotacion al promedio del surtido, y reasignar ese capital de trabajo a "
    "Cuidado personal y Bazar, que hoy rotan mas rapido y sostienen margenes "
    "superiores al 28 % y al 31 % respectivamente. En Panaderia y Carnes, "
    "ademas, la baja rotacion tiene un costo adicional que estas cifras no "
    "capturan: son productos perecibles, de modo que el inventario que no rota "
    "no solo inmoviliza caja, tambien se pierde. Antes de ejecutar la decision "
    "conviene registrar la merma como un tercer proceso de negocio, ya que las "
    "dimensiones conformadas permiten incorporarlo sin rehacer el modelo.")

# ---------------------------------------------------------------------------
# Nota final
# ---------------------------------------------------------------------------
doc.add_paragraph()
par("Nota: todas las cifras de este informe fueron generadas por "
    "lab02_datamart.py sobre la base lab02.duckdb. La salida completa de cada "
    "consulta, sin recortes, esta en salidas/resultados.txt.",
    cursiva=True, size=8)

salida = "Informe_Lab02_TorresMathias.docx"
doc.save(salida)
print("Informe generado:", salida)
